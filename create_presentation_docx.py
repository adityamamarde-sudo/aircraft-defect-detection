import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

doc = Document()

# Set standard slide-report margins (0.75 inch)
for section in doc.sections:
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

# Primary color palette
NAVY_PRIMARY = RGBColor(0, 43, 73)       # #002B49
BLUE_ACCENT = RGBColor(0, 102, 204)     # #0066CC
TEXT_DARK = RGBColor(34, 34, 34)        # Charcoal
TEXT_GRAY = RGBColor(100, 100, 100)

def set_cell_background(cell, hex_color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tc_pr.append(shd)

def add_slide_heading(title_text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(title_text)
    run.font.name = 'Calibri'
    run.font.size = Pt(15)
    run.font.bold = True
    run.font.color.rgb = NAVY_PRIMARY

def add_bullet(text, label=""):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    if label:
        r_lbl = p.add_run(label)
        r_lbl.font.name = 'Calibri'
        r_lbl.font.size = Pt(11)
        r_lbl.font.bold = True
        r_lbl.font.color.rgb = TEXT_DARK
    r_txt = p.add_run(text)
    r_txt.font.name = 'Calibri'
    r_txt.font.size = Pt(11)
    r_txt.font.color.rgb = TEXT_DARK

# ==============================================================================
# SLIDE 1: TITLE & STUDENT DIRECTORY
# ==============================================================================
p_inst = doc.add_paragraph()
p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_inst.paragraph_format.space_after = Pt(2)

r1 = p_inst.add_run("OBJECT ORIENTED PROGRAMMING LAB: IC2305\n")
r1.font.name = 'Calibri'
r1.font.size = Pt(15)
r1.font.bold = True
r1.font.color.rgb = NAVY_PRIMARY

r2 = p_inst.add_run("BRACT’s, Vishwakarma Institute of Technology, Pune\n")
r2.font.name = 'Calibri'
r2.font.size = Pt(12)
r2.font.bold = True

r3 = p_inst.add_run("(An Autonomous Institute affiliated to Savitribai Phule Pune University)\n(NBA and NAAC accredited, ISO 9001:2015 certified)\n")
r3.font.name = 'Calibri'
r3.font.size = Pt(10)
r3.font.color.rgb = TEXT_GRAY

p_topic = doc.add_paragraph()
p_topic.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_topic.paragraph_format.space_before = Pt(8)
p_topic.paragraph_format.space_after = Pt(12)

r_topic = p_topic.add_run("1) Topic: Aircraft Surface Defect Detection System using OOP Concepts")
r_topic.font.name = 'Calibri'
r_topic.font.size = Pt(15)
r_topic.font.bold = True
r_topic.font.color.rgb = BLUE_ACCENT

# Student details table (Strict Order: Roll 03, 16, 10)
t1 = doc.add_table(rows=4, cols=5)
t1.alignment = WD_TABLE_ALIGNMENT.CENTER
headers1 = ["Sr. No.", "Class & Div.", "Roll No.", "P.R. No.", "Name of Student"]
col_widths1 = [Inches(0.8), Inches(1.2), Inches(1.0), Inches(1.6), Inches(2.4)]

for i, h in enumerate(headers1):
    cell = t1.cell(0, i)
    cell.text = h
    set_cell_background(cell, "002B49")
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in p.runs:
        run.font.name = 'Calibri'
        run.font.size = Pt(10.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)

data1 = [
    ["1", "IC-A", "03", "1251140087", "Aditya Mamarde"],
    ["2", "IC-A", "16", "1251140212", "Aryan Gham"],
    ["3", "IC-A", "10", "1251140083", "Anived Thulkar"]
]

for r_idx, r_data in enumerate(data1, start=1):
    bg = "F8F9FA" if r_idx % 2 == 0 else "FFFFFF"
    for c_idx, val in enumerate(r_data):
        cell = t1.cell(r_idx, c_idx)
        cell.text = val
        set_cell_background(cell, bg)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx < 4 else WD_ALIGN_PARAGRAPH.LEFT
        for run in p.runs:
            run.font.name = 'Calibri'
            run.font.size = Pt(10)

for row in t1.rows:
    for idx, width in enumerate(col_widths1):
        row.cells[idx].width = width

p_foot = doc.add_paragraph()
p_foot.paragraph_format.space_before = Pt(8)
p_foot.paragraph_format.space_after = Pt(12)
rf1 = p_foot.add_run("Class: IC-A   |   GROUP NO. 1\n")
rf1.font.bold = True
rf2 = p_foot.add_run("Faculty Guide: PROF. SANIKA PATANKAR")
rf2.font.bold = True
rf2.font.color.rgb = RGBColor(180, 0, 0)

# ==============================================================================
# SLIDE 2: ABSTRACT
# ==============================================================================
add_slide_heading("2) Abstract")
add_bullet("Replaces slow manual inspections with an automated AI computer vision pipeline.", "Automated Surface Inspection: ")
add_bullet("Scans high-resolution aircraft surfaces down to the pixel level to detect minute dents, cracks, and scratches.", "Pixel-Level Detection: ")
add_bullet("Employs Faster R-CNN with ResNet-50 FPN for real-time bounding box localization.", "Deep Learning Model: ")
add_bullet("Provides an interactive slider to filter defects dynamically based on severity and confidence.", "Severity Control: ")
add_bullet("Designed with Encapsulation, Abstraction, and clean separation between inference and GUI logic.", "OOP Principles: ")
add_bullet("Renders real-time bounding boxes, defect counters, and integrity flags via Streamlit.", "Interactive Dashboard: ")

# ==============================================================================
# SLIDE 3: OBJECTIVES
# ==============================================================================
add_slide_heading("3) Objectives")
add_bullet("Detect minute aircraft surface defects in real time using deep neural networks.", "Real-Time Flaw Detection: ")
add_bullet("Accurately localize structural flaws including dents, cracks, scratches, and corrosion.", "Precise Localization: ")
add_bullet("Classify surface anomalies into 6 discrete damage categories.", "Damage Classification: ")
add_bullet("Provide dynamic threshold adjustments to filter minor versus critical damage.", "Severity Thresholding: ")
add_bullet("Encapsulate model inference and data preprocessing pipelines into reusable OOP classes.", "OOP Architecture: ")
add_bullet("Display live bounding box overlays, defect counts, and visual status on a GUI dashboard.", "Dashboard Analytics: ")

# ==============================================================================
# SLIDE 4: FLOW OF THE PROGRAM
# ==============================================================================
add_slide_heading("4) Flow of the Program")
flow_items = [
    "[1] Surface Image Input: Upload high-resolution inspection image via Streamlit.",
    "[2] Media Preprocessing: Frame parsing, RGB color conversion, and dimension formatting.",
    "[3] OOP Tensor Transformation: Pipeline normalization and PyTorch tensor conversion.",
    "[4] Hardware Allocation: Dynamic execution routing to GPU/CUDA or CPU device.",
    "[5] Faster R-CNN Inference: Forward pass through ResNet-50 FPN backbone.",
    "[6] Anomaly Localization: Raw bounding box coordinates and classification scores extracted.",
    "[7] Severity Threshold Filtering: User-defined confidence slider isolates actionable damage.",
    "[8] Annotation Overlay: Matplotlib patches draw bounding rectangles and class tags.",
    "[9] Dashboard Aggregation: Dual-column presentation of raw source vs. AI detection output.",
    "[10] System Verification: Live defect count warnings and inspection reports generated."
]
for item in flow_items:
    add_bullet(item)

# ==============================================================================
# SLIDE 5: TECHNOLOGIES AND LIBRARIES USED
# ==============================================================================
add_slide_heading("5) Technologies and Libraries Used")
t5 = doc.add_table(rows=8, cols=3)
t5.alignment = WD_TABLE_ALIGNMENT.CENTER
headers5 = ["Library / Technology", "Purpose in Project", "Where It Is Used"]
col_widths5 = [Inches(2.0), Inches(2.3), Inches(2.7)]

for i, h in enumerate(headers5):
    cell = t5.cell(0, i)
    cell.text = h
    set_cell_background(cell, "002B49")
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in p.runs:
        run.font.name = 'Calibri'
        run.font.size = Pt(10.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)

tech_data = [
    ["Python", "Primary programming language", "Complete system integration and logic"],
    ["PyTorch (torch)", "Deep learning computation engine", "Tensor processing, CUDA execution, forward pass"],
    ["Torchvision", "Computer vision model backbone", "Faster R-CNN architecture (resnet50_fpn), FastRCNNPredictor"],
    ["Streamlit", "Interactive web UI framework", "Dashboard layout, confidence sliders, session state"],
    ["OpenCV / PIL", "Image parsing and manipulation", "Media ingestion, RGB conversion, frame resizing"],
    ["Matplotlib", "Graphical annotation & rendering", "Overlaying bounding box patches and class tags"],
    ["VS Code", "Development environment", "Writing, running, and debugging code scripts"]
]

for r_idx, r_data in enumerate(tech_data, start=1):
    bg = "F8F9FA" if r_idx % 2 == 0 else "FFFFFF"
    for c_idx, val in enumerate(r_data):
        cell = t5.cell(r_idx, c_idx)
        cell.text = val
        set_cell_background(cell, bg)
        p = cell.paragraphs[0]
        for run in p.runs:
            run.font.name = 'Calibri'
            run.font.size = Pt(9.5)

for row in t5.rows:
    for idx, width in enumerate(col_widths5):
        row.cells[idx].width = width

# ==============================================================================
# SLIDE 6: OOP CONCEPTS APPLIED
# ==============================================================================
add_slide_heading("6) Object-Oriented Programming (OOP) Implementation")
add_bullet("Core system components abstracted into discrete entities (DefectDetector, UIController, Visualizer).", "Classes & Objects: ")
add_bullet("Model weights, device logic, and internal thresholds are private/protected within classes.", "Encapsulation: ")
add_bullet("Complex tensor operations and non-maximum suppression are hidden behind clean .predict() APIs.", "Abstraction: ")
add_bullet("Complete decoupling between image transformation, neural inference, and UI visualization.", "Separation of Concerns: ")
add_bullet("Robust try-except blocks manage CUDA memory constraints and frame capture exceptions safely.", "Exception Handling: ")

# ==============================================================================
# SLIDE 7: MODEL ARCHITECTURE & INFERENCE
# ==============================================================================
add_slide_heading("7) AI Model Architecture & Inference Engine")
add_bullet("Faster R-CNN with a ResNet-50 Feature Pyramid Network (FPN) backbone.", "Neural Backbone: ")
add_bullet("Scratch, Dent, Corrosion, Crack, Rivet Defect, and Surface Peel.", "Detected Defect Classes: ")
add_bullet("High-precision localized box estimation accelerated via CUDA GPU pipelines.", "Inference Routine: ")

# ==============================================================================
# SLIDE 8: UI STATE & EVENT MANAGEMENT
# ==============================================================================
add_slide_heading("8) UI Component State & Event Management")
add_bullet("Encapsulates view states (st.session_state.intro_done) to avoid screen reload loops.", "State Encapsulation: ")
add_bullet("Sliders allow inspectors to adjust confidence thresholds (0.1 to 1.0) on the fly.", "Parameter Control: ")
add_bullet("Synchronizes video intro completion signals cleanly with dashboard initialization.", "Event Synchronization: ")

# ==============================================================================
# SLIDE 9: ANALYTICAL RENDERING & DASHBOARD METRICS
# ==============================================================================
add_slide_heading("9) Analytical Rendering & Dashboard Metrics")
add_bullet("Draws localized rectangles around detected flaws using matplotlib.patches.Rectangle.", "Bounding Overlays: ")
add_bullet("Dynamically tags class identity and confidence percentage above flaw boundaries.", "Real-Time Tagging: ")
add_bullet("Aggregates defect counts per class and triggers warning banners when flaws exceed limits.", "Quantitative Metrics: ")

# ==============================================================================
# SLIDE 10: COMPLETE SYSTEM CLASS STRUCTURE
# ==============================================================================
add_slide_heading("10) Complete System Class Structure")
add_bullet("Loads Faster R-CNN, initializes 6-class predictor heads, and maps weights to device.", "Module 1 (AircraftDefectModel): ")
add_bullet("Normalizes input images into standard PyTorch tensors and sets dimensional channels.", "Module 2 (DataPreprocessor): ")
add_bullet("Executes forward inference pass and filters raw boxes using the confidence threshold.", "Module 3 (DefectDetector): ")
add_bullet("Renders dual-column views and aggregates warning banners and integrity metrics.", "Module 4 (DashboardApp): ")

# ==============================================================================
# SLIDE 11: FINAL SYSTEM OUTPUT & DASHBOARD INSPECTION
# ==============================================================================
add_slide_heading("11) Final System Output & Dashboard Inspection")
add_bullet("Left column shows uploaded raw image; right column shows localized defect detections.", "Dual-Column Inspection: ")
add_bullet("Moving the confidence slider updates bounding boxes and metrics instantaneously.", "Interactive Re-Filtering: ")
add_bullet("Issues defect count alerts (e.g., 'Detected 3 defect(s) above 55% threshold').", "Quantitative Feedback: ")
add_bullet("Outputs a verified pass status when no surface anomalies exceed threshold boundaries.", "Zero-Defect Verification: ")

# ==============================================================================
# SLIDE 12: CONCLUSION
# ==============================================================================
add_slide_heading("12) Conclusion")
add_bullet("Eliminates dangerous and slow manual inspections using automated AI computer vision.", "Automated Inspection: ")
add_bullet("OOP design ensures high reusability, maintainability, and clean decoupling.", "Robust OOP Foundation: ")
add_bullet("Faster R-CNN achieves reliable real-time localization across 6 structural defect classes.", "High Detection Accuracy: ")
add_bullet("Streamlit provides an accessible, responsive dashboard for aerospace maintenance technicians.", "Technician-Centric GUI: ")
add_bullet("Configurable thresholding allows custom sensitivity for minor versus severe damages.", "Severity Calibration: ")

# ==============================================================================
# SLIDE 13: FUTURE SCOPE
# ==============================================================================
add_slide_heading("13) Future Scope")
add_bullet("Quantizing model weights (ONNX/TensorRT) for real-time edge processing on inspection drones.", "Drone Edge Deployment: ")
add_bullet("Projecting 2D defect coordinates directly onto 3D aircraft CAD models for spatial damage maps.", "3D Defect CAD Mapping: ")
add_bullet("Connecting defect logs and severity metrics to enterprise SQL/JDBC maintenance databases.", "Database Backend: ")
add_bullet("Extending inference across multi-camera streams to scan entire fuselage sections continuously.", "Multi-Camera Streaming: ")
add_bullet("Retraining detection heads with larger datasets of obscure micro-cracks.", "Model Refinement: ")
add_bullet("Deploying containerized dashboards on cloud clusters for remote hangar inspection access.", "Cloud Deployment: ")

# ==============================================================================
# SLIDE 14: CLOSING REMARKS
# ==============================================================================
add_slide_heading("14) Thank You")
p_end = doc.add_paragraph()
p_end.paragraph_format.space_before = Pt(6)

rend1 = p_end.add_run("THANK YOU!\n")
rend1.font.name = 'Calibri'
rend1.font.size = Pt(20)
rend1.font.bold = True
rend1.font.color.rgb = NAVY_PRIMARY

rend2 = p_end.add_run("Thank you for your time and attention. Any Questions?\n\n")
rend2.font.name = 'Calibri'
rend2.font.size = Pt(12)

rend3 = p_end.add_run("Built with Python, PyTorch, Torchvision, Streamlit & Matplotlib using OOP Concepts.\nAutomated AI Inspection for Next-Generation Aerospace Safety.")
rend3.font.name = 'Calibri'
rend3.font.size = Pt(10.5)
rend3.font.italic = True
rend3.font.color.rgb = TEXT_GRAY

# Save to .docx
output_path = "Aircraft_Defect_Detection_Presentation_Report.docx"
doc.save(output_path)
print(f"Successfully generated '{output_path}'!")